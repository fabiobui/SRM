"""Feedback dai test sul portale fornitori.

Pochi test mirati: scheda attributi operativi solo per i subappaltatori,
"Salva" che resta sulla scheda, export della dashboard con filtri in AND,
note operative nel payload, stato dei documenti che segue la scadenza,
descrizioni dei servizi IGEAM.
"""

import json
from datetime import timedelta
from io import BytesIO

import openpyxl
import pytest
from django.contrib import admin
from django.core.management import call_command
from django.test import RequestFactory
from django.urls import reverse
from django.utils import timezone

from vendor_management_system.portal.forms import (
    VendorServiceAddRequestForm,
    available_service_types_for_vendor,
)
from vendor_management_system.portal.tests.factories import (
    CategoryFactory,
    DocumentFactory,
    ServiceSetFactory,
    ServiceTypeFactory,
    VendorServiceFactory,
    VendorUserFactory,
)
from vendor_management_system.users.models import User
from vendor_management_system.vendors.admin import (
    VendorAdmin,
    VendorOperationalAttributesInline,
)
from vendor_management_system.vendors.models import ServiceType, Vendor
from vendor_management_system.vendors.tests.factories import VendorFactory

pytestmark = pytest.mark.django_db

PASSWORD = "test-pass-feedback-portale"


@pytest.fixture
def gestore():
    return User.objects.create_superuser(
        email="gestore-feedback@example.invalid",
        password=PASSWORD,
        role="admin",
    )


def _inlines(vendor, gestore):
    request = RequestFactory().get("/")
    request.user = gestore
    model_admin = VendorAdmin(Vendor, admin.site)
    return model_admin.get_inline_instances(request, vendor)


# --- scheda attributi operativi -------------------------------------------


def test_attributi_operativi_visibili_solo_ai_subappaltatori(gestore):
    sub = VendorFactory(vendor_type="Subappaltatore")
    altro = VendorFactory(vendor_type="Consulente")

    assert any(
        isinstance(i, VendorOperationalAttributesInline)
        for i in _inlines(sub, gestore)
    )
    assert not any(
        isinstance(i, VendorOperationalAttributesInline)
        for i in _inlines(altro, gestore)
    )


def test_attributi_operativi_con_classificazione_subappaltatore(gestore):
    principale = VendorFactory(
        vendor_type="Consulente",
        category=CategoryFactory(name="SUBAPPALTATORE"),
    )
    aggiuntiva = VendorFactory(vendor_type="Consulente")
    aggiuntiva.additional_categories.add(
        CategoryFactory(name="Subappaltatore")
    )

    for vendor in (principale, aggiuntiva):
        assert vendor.is_subappaltatore
        assert any(
            isinstance(i, VendorOperationalAttributesInline)
            for i in _inlines(vendor, gestore)
        )


def test_menu_portale_con_classificazione_aggiuntiva_subappaltatore(client):
    utente = VendorUserFactory(
        password=PASSWORD, vendor__vendor_type="Consulente"
    )
    utente.vendor.additional_categories.add(
        CategoryFactory(name="SUBAPPALTATORE")
    )
    client.login(email=utente.email, password=PASSWORD)

    pagina = client.get(reverse("portal:dashboard")).content.decode()

    assert reverse("portal:my-operational-attributes") in pagina


def test_richiesta_servizio_mostra_la_descrizione(client):
    utente = VendorUserFactory(password=PASSWORD)
    con = ServiceTypeFactory(
        parent=ServiceTypeFactory(),
        name="ATEX",
        description="Valutazione del rischio esplosione.",
    )
    ServiceTypeFactory(parent=con.parent, name="Senza descrizione")
    client.login(email=utente.email, password=PASSWORD)

    pagina = client.get(reverse("portal:my-service-add")).content.decode()

    # la descrizione viaggia nella mappa letta dall'icona info; il servizio
    # senza descrizione resta in tendina ma non ha voce nella mappa
    assert "Senza descrizione" in pagina
    mappa = json.loads(
        pagina.split(
            'id="service-descriptions-data" type="application/json">'
        )[1].split("</script>")[0]
    )
    assert list(mappa.values()) == ["Valutazione del rischio esplosione."]


def _servizio_in_set(nome, categoria=None):
    servizio = ServiceTypeFactory(parent=ServiceTypeFactory(), name=nome)
    ServiceSetFactory(category=categoria, is_active=True).service_types.add(
        servizio
    )
    return servizio


def test_servizi_proponibili_set_della_classificazione_e_set_senza():
    sub = CategoryFactory(name="SUBAPPALTATORE")
    mdl = CategoryFactory(name="MDL")
    del_fornitore = _servizio_in_set("Di classificazione", sub)
    per_tutti = _servizio_in_set("Per tutti", None)
    altra_classificazione = _servizio_in_set("Di altra classificazione", mdl)
    fuori_dai_set = ServiceTypeFactory(
        parent=ServiceTypeFactory(), name="Fuori dai set"
    )
    vendor = VendorFactory(category=sub)

    proponibili = set(available_service_types_for_vendor(vendor))

    assert proponibili == {del_fornitore, per_tutti}
    assert altra_classificazione not in proponibili
    assert fuori_dai_set not in proponibili


def test_tendina_servizi_prima_la_classificazione_poi_quelli_per_tutti():
    sub = CategoryFactory(name="SUBAPPALTATORE")
    # "Per tutti" ha un nome e una sezione che verrebbero prima per ordine
    per_tutti = _servizio_in_set("AAA per tutti", None)
    di_classificazione = _servizio_in_set("ZZZ di classificazione", sub)
    vendor = VendorFactory(category=sub)

    form = VendorServiceAddRequestForm(vendor)

    ordine = [
        pk
        for _sezione, voci in form.fields["service_type"].choices
        for pk, _nome in voci
    ]
    assert ordine == [str(di_classificazione.pk), str(per_tutti.pk)]
    assert form.service_priority_ids == [str(di_classificazione.pk)]


def test_servizi_proponibili_senza_classificazione_solo_set_senza():
    per_tutti = _servizio_in_set("Per tutti", None)
    _servizio_in_set("Di classificazione", CategoryFactory(name="MDL"))
    vendor = VendorFactory(category=None)

    assert set(available_service_types_for_vendor(vendor)) == {per_tutti}


def test_menu_portale_attributi_operativi_solo_subappaltatori(client):
    sub = VendorUserFactory(
        password=PASSWORD, vendor__vendor_type="Subappaltatore"
    )
    client.login(email=sub.email, password=PASSWORD)
    url = reverse("portal:my-operational-attributes")
    assert url in client.get(reverse("portal:dashboard")).content.decode()

    altro = VendorUserFactory(
        password=PASSWORD, vendor__vendor_type="Consulente"
    )
    client.login(email=altro.email, password=PASSWORD)
    pagina = client.get(reverse("portal:dashboard")).content.decode()
    assert url not in pagina


def test_navbar_portale_non_mostra_il_codice_fornitore(client):
    utente = VendorUserFactory(password=PASSWORD)
    client.login(email=utente.email, password=PASSWORD)

    pagina = client.get(reverse("portal:dashboard")).content.decode()

    assert utente.vendor.vendor_code not in pagina


# --- "Salva" resta sulla scheda ---------------------------------------------


def test_salva_equivale_a_salva_e_continua(gestore):
    request = RequestFactory().post("/", {"_save": "Salva"})
    VendorAdmin(Vendor, admin.site)._salva_resta_sulla_scheda(request)
    assert "_continue" in request.POST


def test_la_scheda_fornitore_non_ha_salva_e_continua(client, gestore):
    vendor = VendorFactory()
    client.login(email=gestore.email, password=PASSWORD)

    html = client.get(
        reverse("admin:vendors_vendor_change", args=[vendor.pk])
    ).content.decode()

    assert 'name="_save"' in html
    assert 'name="_continue"' not in html


# --- dashboard ---


def _nomi_export(risposta):
    foglio = openpyxl.load_workbook(BytesIO(risposta.content)).active
    return [
        foglio.cell(row=r, column=1).value
        for r in range(2, foglio.max_row + 1)
        if foglio.cell(row=r, column=1).value
    ]


def test_export_con_piu_servizi_richiede_tutti(client, gestore):
    s1 = ServiceTypeFactory(parent=ServiceTypeFactory(), name="ATEX")
    s2 = ServiceTypeFactory(parent=ServiceTypeFactory(), name="DVR")
    entrambi = VendorFactory(name="Ha entrambi")
    uno = VendorFactory(name="Ha solo ATEX")
    for svc in (s1, s2):
        VendorServiceFactory(vendor=entrambi, service_type=svc)
    VendorServiceFactory(vendor=uno, service_type=s1)
    client.login(email=gestore.email, password=PASSWORD)

    risposta = client.get(
        reverse("export-vendors-excel"), {"services": "ATEX,DVR"}
    )

    assert _nomi_export(risposta) == ["Ha entrambi"]


def test_payload_dashboard_include_note_operative_e_stato_contrattuale(
    client, gestore
):
    VendorFactory(name="Con note", contact_details="Referente: Mario Rossi")
    client.login(email=gestore.email, password=PASSWORD)

    risposta = client.get(reverse("vendor-dashboard"))

    dati = json.loads(risposta.context["vendors_data_json"])
    riga = next(v for v in dati if v["name"] == "Con note")
    assert riga["operational_notes"] == "Referente: Mario Rossi"
    assert "contractual_status" in riga


# --- stato documenti e scadenza ---


def test_documento_scaduto_torna_approvato_se_la_scadenza_va_nel_futuro():
    oggi = timezone.now().date()
    doc = DocumentFactory(
        status="APPROVED", expiry_date=oggi - timedelta(days=1)
    )
    doc.reviewed_at = timezone.now()
    doc.save()
    doc.refresh_from_db()
    assert doc.status == "EXPIRED"

    doc.expiry_date = oggi + timedelta(days=90)
    doc.save()

    doc.refresh_from_db()
    assert doc.status == "APPROVED"


def test_documento_scaduto_mai_approvato_non_diventa_approvato():
    oggi = timezone.now().date()
    doc = DocumentFactory(
        status="UPLOADED", expiry_date=oggi - timedelta(days=1)
    )
    assert doc.status == "EXPIRED"

    doc.expiry_date = oggi + timedelta(days=30)
    doc.save()

    doc.refresh_from_db()
    assert doc.status in ("UPLOADED", "PENDING")


def test_job_sync_segna_scaduti_i_documenti_con_scadenza_passata():
    from vendor_management_system.vendors.qualification_sync import sync_all

    doc = DocumentFactory(status="APPROVED")
    type(doc).objects.filter(pk=doc.pk).update(
        expiry_date=timezone.now().date() - timedelta(days=2)
    )

    riepilogo = sync_all()

    doc.refresh_from_db()
    assert doc.status == "EXPIRED"
    assert riepilogo["expired_documents"] == 1


# --- descrizioni servizi IGEAM ---


def test_seed_servizi_salva_le_descrizioni_senza_sovrascrivere_le_modifiche():
    call_command("seed_service_sets", "--create-only")
    servizio = ServiceType.objects.get(code="SERV-0016")
    assert servizio.description.startswith("Valutazione e gestione")

    servizio.description = "Modificata da admin"
    servizio.save()
    call_command("seed_service_sets", "--create-only")

    servizio.refresh_from_db()
    assert servizio.description == "Modificata da admin"
