"""Dashboard e export basati sulla zona di competenza.

Prima di questo lavoro i grafici "Fornitori per Regione" e "Province"
contavano la SEDE del fornitore (`Address.region` / `state_province`,
stringhe libere). Ora contano la copertura territoriale dichiarata.
"""

import json
from io import BytesIO

import pytest
from django.urls import reverse

from vendor_management_system.users.models import User
from vendor_management_system.vendors import competence
from vendor_management_system.vendors.models import Province
from vendor_management_system.vendors.tests.factories import (
    AddressFactory,
    VendorFactory,
)

PASSWORD = "test-pass-dashboard-competence"


def _login(client, etichetta="dashboard-competence"):
    utente = User.objects.create_superuser(
        email=f"{etichetta}@example.invalid",
        password=PASSWORD,
        role="admin",
    )
    client.login(email=utente.email, password=PASSWORD)
    return utente


def _chart_data(risposta):
    return json.loads(risposta.context["chart_data_json"])


def _vendors_data(risposta):
    return json.loads(risposta.context["vendors_data_json"])


def _province_di(codice_regione):
    return list(
        Province.objects.filter(region__code=codice_regione).values_list(
            "code", flat=True
        )
    )


# --------------------------------------------------------------------------
# Grafici
# --------------------------------------------------------------------------


@pytest.mark.django_db
def test_le_regioni_contano_la_competenza_non_la_sede(
    geografia_italia, client
):
    vendor = VendorFactory(address=AddressFactory(region="Lazio"))
    competence.set_vendor_competence(
        vendor, province_codes=["MI"], country_codes=[]
    )
    _login(client)

    dati = _chart_data(client.get(reverse("vendor-dashboard")))

    per_regione = {r["region"]: r["count"] for r in dati["by_region"]}
    assert per_regione.get("Lombardia") == 1
    assert "Lazio" not in per_regione


@pytest.mark.django_db
def test_un_fornitore_conta_in_ogni_regione_coperta(geografia_italia, client):
    vendor = VendorFactory()
    competence.set_vendor_competence(
        vendor, province_codes=["MI", "RM", "NA", "TO", "BO"], country_codes=[]
    )
    _login(client)

    dati = _chart_data(client.get(reverse("vendor-dashboard")))

    barre = [r for r in dati["by_region"] if r["code"]]
    assert len(barre) == 5
    assert sum(r["count"] for r in barre) == 5


@pytest.mark.django_db
def test_una_regione_intera_vale_uno_non_dodici(geografia_italia, client):
    vendor = VendorFactory()
    competence.set_vendor_competence(
        vendor, province_codes=_province_di("LOM"), country_codes=[]
    )
    _login(client)

    dati = _chart_data(client.get(reverse("vendor-dashboard")))

    lombardia = next(r for r in dati["by_region"] if r["code"] == "LOM")
    assert lombardia["count"] == 1


@pytest.mark.django_db
def test_le_province_sono_calcolate_lato_server(geografia_italia, client):
    """Prima non lo erano: il JavaScript le ricalcolava sempre nel
    browser, perche' la view non produceva mai `by_province`."""
    vendor = VendorFactory()
    competence.set_vendor_competence(
        vendor, province_codes=["MI"], country_codes=[]
    )
    _login(client)

    dati = _chart_data(client.get(reverse("vendor-dashboard")))

    assert dati["by_province"] == [
        {
            "province": "Milano",
            "code": "MI",
            "region": "LOM",
            "region_name": "Lombardia",
            "count": 1,
        }
    ]


@pytest.mark.django_db
def test_bucket_non_specificato(geografia_italia, client):
    VendorFactory()
    coperto = VendorFactory()
    competence.set_vendor_competence(
        coperto, province_codes=["MI"], country_codes=[]
    )
    _login(client)

    dati = _chart_data(client.get(reverse("vendor-dashboard")))

    bucket = next(r for r in dati["by_region"] if r["code"] is None)
    assert bucket["region"] == "Non specificato"
    assert bucket["count"] == 1
    assert dati["competence_meta"]["vendors_without_competence"] == 1


@pytest.mark.django_db
def test_nazioni_estere_aggregate(geografia_italia, client):
    vendor = VendorFactory()
    competence.set_vendor_competence(
        vendor, province_codes=[], country_codes=["FR"]
    )
    _login(client)

    dati = _chart_data(client.get(reverse("vendor-dashboard")))

    assert dati["by_competence_country"] == [
        {"country": "Francia", "code": "FR", "count": 1}
    ]
    # Ha una competenza, quindi non finisce nel bucket dei non
    # specificati: il JavaScript applica la stessa definizione.
    assert dati["competence_meta"]["vendors_without_competence"] == 0


# --------------------------------------------------------------------------
# Payload della tabella
# --------------------------------------------------------------------------


@pytest.mark.django_db
def test_il_payload_porta_i_codici_di_competenza(geografia_italia, client):
    vendor = VendorFactory()
    competence.set_vendor_competence(
        vendor, province_codes=["MI"], country_codes=["FR"]
    )
    _login(client)

    dati = _vendors_data(client.get(reverse("vendor-dashboard")))

    assert dati[0]["competence"] == {
        "provinces": ["MI"],
        "regions": ["LOM"],
        "countries": ["FR"],
    }


@pytest.mark.django_db
def test_copertura_totale_usa_la_sentinella(geografia_italia, client):
    """Senza, un fornitore su tutta Italia peserebbe ~0,7 KB nel payload
    di una pagina che non passa da GZipMiddleware."""
    vendor = VendorFactory()
    competence.set_vendor_competence(
        vendor,
        province_codes=list(Province.objects.values_list("code", flat=True)),
        country_codes=[],
    )
    _login(client)

    dati = _vendors_data(client.get(reverse("vendor-dashboard")))

    assert dati[0]["competence"]["provinces"] == ["*"]
    assert dati[0]["competence"]["regions"] == ["*"]


@pytest.mark.django_db
def test_il_titolo_del_grafico_e_regioni(geografia_italia, client):
    _login(client)

    contenuto = client.get(reverse("vendor-dashboard")).content.decode()

    assert '<h5 class="mb-1">Regioni</h5>' in contenuto
    assert "Fornitori per Regione" not in contenuto
    assert "conteggiato in ognuna" in contenuto


# --------------------------------------------------------------------------
# Export Excel
# --------------------------------------------------------------------------


@pytest.mark.django_db
def test_export_filtra_per_provincia_di_competenza(geografia_italia, client):
    coperto = VendorFactory(name="Copre Milano")
    escluso = VendorFactory(name="Copre Roma")
    competence.set_vendor_competence(
        coperto, province_codes=["MI"], country_codes=[]
    )
    competence.set_vendor_competence(
        escluso, province_codes=["RM"], country_codes=[]
    )
    _login(client)

    risposta = client.get(
        reverse("export-vendors-excel"), {"comp_provinces": "MI"}
    )

    nomi = _nomi_export(risposta)
    assert "Copre Milano" in nomi
    assert "Copre Roma" not in nomi


@pytest.mark.django_db
def test_export_filtra_per_regione_di_competenza(geografia_italia, client):
    coperto = VendorFactory(name="Lombardo")
    escluso = VendorFactory(name="Laziale")
    competence.set_vendor_competence(
        coperto, province_codes=["MI"], country_codes=[]
    )
    competence.set_vendor_competence(
        escluso, province_codes=["RM"], country_codes=[]
    )
    _login(client)

    risposta = client.get(
        reverse("export-vendors-excel"), {"comp_regions": "LOM"}
    )

    nomi = _nomi_export(risposta)
    assert nomi == ["Lombardo"]


@pytest.mark.django_db
def test_export_filtra_per_nazione_di_competenza(geografia_italia, client):
    francese = VendorFactory(name="Francese")
    italiano = VendorFactory(name="Italiano")
    competence.set_vendor_competence(
        francese, province_codes=[], country_codes=["FR"]
    )
    competence.set_vendor_competence(
        italiano, province_codes=["MI"], country_codes=[]
    )
    _login(client)

    risposta = client.get(
        reverse("export-vendors-excel"), {"comp_countries": "FR"}
    )

    assert _nomi_export(risposta) == ["Francese"]


@pytest.mark.django_db
def test_export_non_duplica_le_righe(geografia_italia, client):
    vendor = VendorFactory(name="Tutta la Lombardia")
    competence.set_vendor_competence(
        vendor, province_codes=_province_di("LOM"), country_codes=[]
    )
    _login(client)

    risposta = client.get(
        reverse("export-vendors-excel"), {"comp_regions": "LOM"}
    )

    assert _nomi_export(risposta) == ["Tutta la Lombardia"]


@pytest.mark.django_db
def test_export_ha_le_colonne_di_competenza(geografia_italia, client):
    vendor = VendorFactory()
    competence.set_vendor_competence(
        vendor, province_codes=["MI"], country_codes=[]
    )
    _login(client)

    foglio = _foglio_export(client.get(reverse("export-vendors-excel")))

    intestazioni = [cella.value for cella in foglio[1]]
    assert intestazioni[-2:] == [
        "Zone di Competenza",
        "Province Competenza",
    ]
    assert "Milano" in foglio.cell(row=2, column=8).value
    assert foglio.cell(row=2, column=9).value == "MI"


@pytest.mark.django_db
def test_i_vecchi_parametri_sede_non_filtrano_piu(geografia_italia, client):
    """`regions`/`provinces` filtravano per nome sulla sede e non esistono
    piu': se il JavaScript li inviasse ancora, l'export ignorerebbe in
    silenzio i filtri geografici."""
    coperto = VendorFactory(name="Copre Milano")
    altro = VendorFactory(name="Copre Roma")
    competence.set_vendor_competence(
        coperto, province_codes=["MI"], country_codes=[]
    )
    competence.set_vendor_competence(
        altro, province_codes=["RM"], country_codes=[]
    )
    _login(client)

    risposta = client.get(
        reverse("export-vendors-excel"), {"provinces": "Milano"}
    )

    assert sorted(_nomi_export(risposta)) == ["Copre Milano", "Copre Roma"]


def _foglio_export(risposta):
    import openpyxl

    return openpyxl.load_workbook(BytesIO(risposta.content)).active


def _nomi_export(risposta):
    foglio = _foglio_export(risposta)
    return [
        foglio.cell(row=riga, column=1).value
        for riga in range(2, foglio.max_row + 1)
        if foglio.cell(row=riga, column=1).value
    ]
